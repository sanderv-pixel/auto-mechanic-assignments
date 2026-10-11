#!/usr/bin/env python3
"""Turn a simple Markdown file into a professionally styled Word document (.docx).

Supported input (one line = one paragraph, blank lines are ignored):

  ---                        optional front matter
  title: Brake Systems       settings: title, subtitle, kicker, theme, lang, cover, toc,
  subtitle: Assignment 3               toc_title, font, size, spacing
  kicker: Research Paper     small capitals label above the title (optional)
  theme: navy                navy (default), graphite, or plain (black and white, for strict formats)
  lang: en-US                en-US (Letter) or en-GB (A4); sets spell-check language too
  cover: true                separate cover page (default: true when a title is set)
  toc: false                 contents page after the cover, with real page numbers when LibreOffice exists
  font: Arial                size: 11   spacing: 1.5   ("12pt" and "1,5" are understood)
  Name: [Name]               any other key is a detail on the cover / header block,
  Class: AUTO 110            in the order you write it
  ---
  # Heading 1 / ## Heading 2 / ### Heading 3   (deeper levels become Heading 3)
  **1. Question text (2 pts)**   a whole bold line starting with a number, "Question" or "Task" is
                                 a question box; a bold line starting with a), b) is a sub-question
  Normal paragraph with **bold**, *italic*, ***both*** and `code` (shown as plain text).
  - bullet            (indent two spaces for a sub-bullet)
  1. numbered item    (also a), b) ... kept literally with a hanging indent)
  > callout           shaded note box: given data, formulas, safety notes
  | table | header |  (Markdown pipe table, first row = header)
  ![Caption text](fig1.png){width=9}   image, path relative to this file, width in cm optional
  ![Figure 2: Old and new pads](placeholder "Photo of the pads you removed next to the new ones"){height=7}
                      dashed "Insert picture here" box for a picture the student adds by hand;
                      an image that can't be found or read also becomes such a box
  \\pagebreak

Lines under a "References" / "Works Cited" / "Bibliography" heading get an APA-style hanging
indent. URLs and fill-in blanks (___) are never treated as formatting.

Safety nets: every em dash and en dash is removed (body, cover, header, metadata, image names)
and reported, and the finished file carries no traces of the script that made it: template ids,
preview thumbnail and leftover parts are removed, the file properties name the student and give
real counts, and the package is written the way Word writes it (see docx_finalize.py).

Usage:
  python3 build_docx.py answers.md -o Assignment.docx [--theme navy] [--font Arial] [--size 11]
                        [--spacing 1.5] [--no-render]
"""

from __future__ import annotations

import argparse
import math
import re
import shutil
import sys
import tempfile
from dataclasses import dataclass, field
from pathlib import Path

try:
    from docx import Document
    from docx.document import Document as WordDocument
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt
    from lxml import etree
except ImportError:
    sys.exit("python-docx is missing. Install it with: pip install python-docx")

sys.path.insert(0, str(Path(__file__).resolve().parent))
import docx_finalize  # noqa: E402
import docx_style as ds  # noqa: E402
import office_render  # noqa: E402

FIGURE_DASH, EN_DASH, EM_DASH, BAR = chr(0x2012), chr(0x2013), chr(0x2014), chr(0x2015)
ALL_DASHES = FIGURE_DASH + EN_DASH + EM_DASH + BAR

LANG_CODES = {"en-us": "en-US", "en": "en-US", "us": "en-US", "american": "en-US", "en-ca": "en-US",
              "en-gb": "en-GB", "uk": "en-GB", "gb": "en-GB", "british": "en-GB", "en-au": "en-GB",
              "en-nz": "en-GB", "en-ie": "en-GB", "en-za": "en-GB", "en-in": "en-GB"}
SETTING_KEYS = {"title", "subtitle", "kicker", "theme", "lang", "cover", "toc", "toc_title", "font", "size", "spacing"}
TOC_LABEL = "Contents"
NAME_KEYS = {"name", "student", "student name", "author"}
COURSE_KEYS = {"class", "course", "course name", "subject", "module", "unit"}
REFERENCES = re.compile(r"^(references|works cited|bibliography|sources|reference list)\s*:?$", re.IGNORECASE)
HANGING_INDENT_CM = 1.27
DEFAULT_IMAGE_WIDTH_CM = 14.0
MAX_IMAGE_WIDTH_CM = 16.0
MAX_IMAGE_HEIGHT_CM = 18.0  # leaves room for the question and caption on the same page
PLACEHOLDER_SIZE_CM = (12.0, 7.0)  # default width, height of a picture placeholder box
PLACEHOLDER_PATHS = {"", "placeholder", "photo", "picture", "image"}
HEADER_TITLE_MAX = 60
WORD_IMAGE_TYPES = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tif", ".tiff"}
NBSP = chr(0xA0)
CONTROL_CHARS = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]")  # not allowed in Word XML
PAGEBREAKS = {"\\pagebreak", "<!-- pagebreak -->", "[pagebreak]"}
METADATA_MAX = 255  # Word's limit for the author field


@dataclass(frozen=True)
class Settings:
    font: str = "Calibri"
    size: float = 11.0
    spacing: float = 1.15
    lang: str = "en-US"
    cover: bool | None = None
    toc: bool = False
    title: str = ""
    subtitle: str = ""
    kicker: str = ""
    theme: str = "navy"
    toc_title: str = ""


@dataclass
class Report:
    dashes_fixed: int = 0
    image_problems: list[str] = field(default_factory=list)
    pictures: list[str] = field(default_factory=list)  # picture placeholders the student has to fill
    notes: list[str] = field(default_factory=list)


# ---------------------------------------------------------------- dash cleanup

def strip_dashes(text: str, report: Report) -> str:
    """Remove em/en dashes (and '--') in a way that keeps the sentence readable."""
    double_hyphens = re.findall(r"(?<=\w)--(?=\w)|(?:^|\s)--(?:\s|$)", text)
    found = sum(text.count(char) for char in ALL_DASHES) + len(double_hyphens)
    if not found:
        return text
    report.dashes_fixed += found
    dash = f"[{ALL_DASHES}]"
    text = re.sub(rf"(\d)\s*{dash}\s*(\d)", r"\1-\2", text)  # number ranges: 10-12
    text = re.sub(rf"^\s*(?:{dash}|--)\s*", "", text)  # dash at the very start
    text = re.sub(rf"\s*(?:[{EM_DASH}{BAR}]|--)\s*$", "", text)  # dash at the very end
    text = re.sub(rf"\s*(?:[{EM_DASH}{BAR}]|--)\s*", ", ", text)  # em dash between words
    text = re.sub(rf"\s+[{EN_DASH}{FIGURE_DASH}]\s+", ", ", text)  # spaced en dash used as a pause
    text = re.sub(rf"[{EN_DASH}{FIGURE_DASH}]", "-", text)  # unspaced leftovers: compound words
    text = re.sub(r"([:;,])\s*,\s*", r"\1 ", text)  # "Note: , x" -> "Note: x"
    text = re.sub(r",\s*([.;:!?)])", r"\1", text)
    return re.sub(r"\(\s*,\s*", "(", text)


# ------------------------------------------------------------- parsing input

def parse_front_matter(lines: list[str]) -> tuple[dict[str, str], list[tuple[str, str]], list[str]]:
    if not lines or lines[0].strip() != "---":
        return {}, [], lines
    settings: dict[str, str] = {}
    details: list[tuple[str, str]] = []
    for index in range(1, len(lines)):
        line = lines[index]
        if line.strip() == "---":
            return settings, details, lines[index + 1:]
        if ":" not in line:
            continue
        key, value = line.split(":", 1)
        key, value = key.strip(), value.strip().strip('"').strip("'")
        if key.lower() in SETTING_KEYS:
            settings[key.lower()] = value
        elif key:
            details.append((key[:1].upper() + key[1:], value))
    return {}, [], lines  # no closing ---, treat everything as body


def to_bool(value: str | None) -> bool | None:
    if value is None:
        return None
    return value.strip().lower() in {"true", "yes", "1", "on"}


def parse_number(value: object, default: float, low: float, high: float) -> float:
    """Read '12', '12pt', '1,5' or 1.5 and keep it in a sane range."""
    match = re.search(r"\d+(?:[.,]\d+)?", str(value)) if value not in (None, "") else None
    if not match:
        return default
    return min(high, max(low, float(match.group().replace(",", "."))))


def lang_code(raw: dict[str, str]) -> str | None:
    """en-US / en-GB, or None when the value is not recognised."""
    return LANG_CODES.get((raw.get("lang") or "en-US").strip().lower().replace("_", "-"))


def build_settings(raw: dict[str, str], args: argparse.Namespace) -> Settings:
    lang = lang_code(raw) or "en-US"
    theme = (args.theme or raw.get("theme") or Settings.theme).strip().lower()
    font = re.sub(f"[{ALL_DASHES}]", "-", (args.font or raw.get("font") or Settings.font).strip())
    return Settings(
        font=font,
        size=parse_number(args.size or raw.get("size"), Settings.size, 8, 20),
        spacing=parse_number(args.spacing or raw.get("spacing"), Settings.spacing, 1, 3),
        lang=lang,
        cover=to_bool(raw.get("cover")),
        toc=bool(to_bool(raw.get("toc"))),
        title=raw.get("title", ""),
        subtitle=raw.get("subtitle", ""),
        kicker=raw.get("kicker", ""),
        theme=theme if theme in ds.THEMES else Settings.theme,
        toc_title=raw.get("toc_title", ""),
    )


BOLD_ITALIC = r"\*\*\*(?=\S)[^*]+?(?<=\S)\*\*\*"
BOLD = r"\*\*(?=\S)(?:[^*]|\*(?!\*))+?(?<=\S)\*\*"  # may contain *italic*
ITALIC_STAR = r"(?<![\w*])\*(?=[^\s*])[^*]*?(?<=[^\s*])\*(?![\w*])"
ITALIC_UNDERSCORE = r"(?<![\w_])_(?=[^\s_])[^_]*?(?<=[^\s_])_(?![\w_])"  # never matches ___ blanks
CODE = r"`[^`]+`"
INLINE = re.compile(f"({BOLD_ITALIC}|{BOLD}|{ITALIC_STAR}|{ITALIC_UNDERSCORE}|{CODE})")
URL = re.compile(r"((?:https?://|www\.)\S+)")


def add_run(paragraph, text: str, bold: bool, italic: bool):
    run = paragraph.add_run(text)
    run.bold = bold or None
    run.italic = italic or None
    return run


def clean_text(text: str, report: Report) -> str:
    """strip_dashes outside URLs; inside a URL only real dash characters become hyphens."""
    urls: list[str] = []

    def hold(match: re.Match) -> str:
        urls.append(match.group(0))
        return f"\ue000{len(urls) - 1}\ue001"

    def restore(match: re.Match) -> str:
        url = urls[int(match.group(1))]
        report.dashes_fixed += sum(url.count(char) for char in ALL_DASHES)
        return re.sub(f"[{ALL_DASHES}]", "-", url)

    return re.sub("\ue000(\\d+)\ue001", restore, strip_dashes(URL.sub(hold, text), report))


def inline_parts(text: str, bold: bool = False, italic: bool = False) -> list[tuple[str, bool, bool]]:
    """Split Markdown into (text, bold, italic) pieces. URLs stay literal: _ in a link is not italics."""
    parts: list[tuple[str, bool, bool]] = []
    for chunk in URL.split(text):
        if URL.fullmatch(chunk):
            parts.append((chunk, bold, italic))
            continue
        for part in INLINE.split(chunk):
            if not part:
                continue
            if re.fullmatch(BOLD_ITALIC, part):
                parts.append((part[3:-3], True, True))
            elif re.fullmatch(BOLD, part):
                parts += inline_parts(part[2:-2], True, italic)
            elif re.fullmatch(CODE, part):
                parts.append((part[1:-1], bold, italic))
            elif re.fullmatch(f"{ITALIC_STAR}|{ITALIC_UNDERSCORE}", part):
                parts.append((part[1:-1], bold, True))
            else:
                parts.append((part, bold, italic))
    return parts


def add_inline(paragraph, text: str, report: Report, bold: bool = False, italic: bool = False) -> list:
    """Add Markdown-formatted text to a paragraph and return the runs it made."""
    return [add_run(paragraph, piece, b, i) for piece, b, i in inline_parts(clean_text(text, report), bold, italic)]


def plain_text(text: str) -> str:
    """The text add_inline would show, without formatting (for metadata, header and contents)."""
    return "".join(piece for piece, _, _ in inline_parts(clean_text(text, Report()))).strip()


def clean_heading(text: str) -> str:
    return re.sub(r"\*+|`|(?<![\w_])_|_(?![\w_])", "", text).strip()


def is_placeholder(value: str) -> bool:
    value = value.strip()
    return not value or "[" in value or set(value) <= {"_", " "}


def detail_value(details: list[tuple[str, str]], keys: set[str]) -> str:
    for key, value in details:
        if key.lower() in keys and not is_placeholder(value):
            return value
    return ""


# ------------------------------------------------------------ document parts

def page_break(doc) -> None:
    """A 1 pt paragraph holding the break, so nothing visible lands at the top of the next page."""
    ds.exact_spacer(doc, 1).add_run().add_break(WD_BREAK.PAGE)


def add_kicker(doc, settings: Settings, report: Report) -> None:
    if settings.kicker:
        add_inline(doc.add_paragraph(style=ds.KICKER_STYLE), clean_heading(settings.kicker), report)


def cover_spacing(settings: Settings, details: list[tuple[str, str]]) -> tuple[float, float]:
    """Space above the title and above the details, so the cover always fits one page."""
    scale = settings.size / Settings.size  # detail rows grow with the body font size
    title_lines = max(1, math.ceil(len(clean_heading(settings.title)) / 24))
    subtitle_lines = math.ceil(len(settings.subtitle) / 48)
    detail_lines = sum(max(1, math.ceil(len(value) * scale / 50)) for _, value in details)
    used = (10 + (20 if settings.kicker else 0) + 39 * title_lines + 14 + 20 * subtitle_lines + 12
            + 15 * scale * detail_lines + 10 * len(details) + 60 + 3 * settings.size)  # safety margin
    free = ds.usable_height_pt(settings.lang) - used
    top = min(150.0, max(24.0, free * 0.42))
    return top, max(24.0, free - top)


def add_cover(doc, settings: Settings, details: list[tuple[str, str]], theme: ds.Theme, report: Report) -> None:
    """Cover page: accent bar, title block in the upper third, details table at the bottom."""
    top, middle = cover_spacing(settings, details)
    total_cm = ds.text_width_cm(settings.lang)
    ds.set_borders(ds.exact_spacer(doc, 4), bottom=(36, theme.primary, 0))
    ds.exact_spacer(doc, top)
    add_kicker(doc, settings, report)
    title = doc.add_paragraph(style="Title")
    ds.single_spaced(title, after=2)
    add_inline(title, settings.title, report)
    rule = ds.exact_spacer(doc, 10)
    rule.paragraph_format.right_indent = Cm(total_cm - 2.5)
    ds.set_borders(rule, bottom=(24, theme.accent, 0))
    if settings.subtitle:
        subtitle = doc.add_paragraph(style="Subtitle")
        ds.single_spaced(subtitle, before=12)
        add_inline(subtitle, settings.subtitle, report)
    ds.exact_spacer(doc, middle)
    if details:
        ds.details_table(doc, details, theme, total_cm, lambda p, text: add_inline(p, text, report))
    page_break(doc)


def add_header_block(doc, settings: Settings, details: list[tuple[str, str]], theme: ds.Theme,
                     report: Report) -> None:
    """Compact top-of-page block for worksheets and short tasks (no separate cover)."""
    add_kicker(doc, settings, report)
    if settings.title:
        add_inline(doc.add_paragraph(style=ds.TITLE_STYLE), settings.title, report)
    if settings.subtitle:
        subtitle = doc.add_paragraph(style=ds.DETAILS_STYLE)
        subtitle.paragraph_format.space_after = Pt(4)
        ds.style_runs(add_inline(subtitle, settings.subtitle, report), size=12, color=theme.accent)
    if details:
        ds.exact_spacer(doc, 6)
        ds.info_bar(doc, details, theme, ds.text_width_cm(settings.lang), lambda p, text: add_inline(p, text, report))
    if settings.title or details:
        ds.exact_spacer(doc, 12)


def collect_headings(lines: list[str], max_level: int = 2) -> list[tuple[int, str]]:
    """(level, text) of every heading up to max_level, exactly as the heading itself will read."""
    entries = []
    for line in lines:
        match = re.match(rf"^(#{{1,{max_level}}})\s+(.*)$", CONTROL_CHARS.sub("", line).strip())
        if match:
            entries.append((len(match.group(1)), plain_text(match.group(2))))
    return entries


def add_toc(doc, settings: Settings, entries: list[tuple[int, str]], pages: dict[int, int] | None,
            report: Report) -> None:
    """A real Word TOC field, pre-filled with the headings and (when known) their page numbers.

    Word refreshes it when the student right-clicks it and picks Update Field.
    """
    heading = doc.add_paragraph(style="TOC Heading")
    add_inline(heading, settings.toc_title or TOC_LABEL, report)
    entries = entries or [(1, "")]
    for index, (level, text) in enumerate(entries):
        paragraph = doc.add_paragraph(style=f"toc {min(level, 2)}")
        if index == 0:
            ds.field_start(paragraph, 'TOC \\o "1-2" \\h \\z \\u')
        paragraph.add_run(text)
        if pages and index in pages:
            paragraph.add_run(f"\t{pages[index]}")
        if index == len(entries) - 1:
            ds.field_end(paragraph)
    page_break(doc)


def add_page_furniture(doc, settings: Settings, details: list[tuple[str, str]], theme: ds.Theme,
                       use_cover: bool) -> None:
    """Running header (title left, student right) and 'Page X of Y' footer; none on the cover."""
    section = doc.sections[0]
    section.different_first_page_header_footer = True
    section.first_page_header.is_linked_to_previous = False
    title = plain_text(settings.title)
    if len(title) > HEADER_TITLE_MAX:
        title = title[:HEADER_TITLE_MAX].rsplit(" ", 1)[0] + "..."
    right = plain_text(detail_value(details, NAME_KEYS) or detail_value(details, COURSE_KEYS))[:HEADER_TITLE_MAX]
    if title or right:
        ds.page_header(section, title, right, theme, ds.text_width_cm(settings.lang))
    ds.page_footer(section.footer, theme)
    if use_cover:
        section.first_page_footer.is_linked_to_previous = False
    else:
        ds.page_footer(section.first_page_footer, theme)


# ------------------------------------------------------------------ body blocks

TABLE_SEPARATOR = re.compile(r"^\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$")
NUMBERED = re.compile(r"^(\d{1,3}[.)]|[a-zA-Z][.)]|[ivx]{1,4}[.)])\s+(.*)$")
IMAGE = re.compile(r'^!\[(?P<caption>.*?)\]\((?P<path>(?:[^()"]|\([^()"]*\))*?)\s*(?:"(?P<hint>[^"]*)")?\s*\)'
                   r'(?:\{(?P<attrs>[^}]*)\})?\s*$')  # paths may hold (1) like "Screenshot (1).png"
LOOSE_IMAGE = re.compile(r"^!\[(?P<caption>.*?)\]\(.*\)")
HEADING = re.compile(r"^(#{1,6})\s+(.*)$")
FULL_BOLD = re.compile(r"^\*\*(?!\*)((?:(?!\*\*).)+?)\*\*$")
QUESTION_START = re.compile(r"^(?:\d{1,3}[.)]\s|\d{1,2}\.\d{1,2}[.)]?\s+[A-Z][a-z]|"  # 1. / 2.1 Explain (not 3.5 L)
                            r"(?i:question|task|problem|exercise|activity|part|q)\s*\d+[a-zA-Z]?\b)")
SUB_START = re.compile(r"^(?:[a-hA-H]|[ivx]{1,4})[.)]\s")
OPTION_MARK = re.compile(r"(?:^|[\s,;])([a-hA-H])([.)])(?=\s)")
POINTS = re.compile(r"\s\((\d+(?:[.,]\d+)?\s*(?:pts?|points?|marks?))\)$", re.IGNORECASE)
NUMERIC_CELL = re.compile(r"^[~<>+\-±]?\s?\d[\d.,]*(?:\s?(?:-|to)\s?\d[\d.,]*)?\s?[^\s\d]{0,8}$")


def split_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def is_options_line(text: str) -> bool:
    """'a) DOT 3, b) DOT 5, ...': at least an a and a b marker, same case and punctuation."""
    marks = OPTION_MARK.findall(text)
    return (len(marks) >= 2 and [letter for letter, _ in marks[:2]] in (["a", "b"], ["A", "B"])
            and marks[0][1] == marks[1][1])


def question_kind(text: str, previous: str) -> str | None:
    """'question', 'sub' or None for a line that is bold from start to end."""
    if QUESTION_START.match(text):
        return "question"
    if SUB_START.match(text):
        return "question" if is_options_line(text) and previous == "question" else "sub"
    return "question" if previous == "question" else None


def add_question(doc, text: str, kind: str, continues: bool, theme: ds.Theme, settings: Settings,
                 report: Report) -> None:
    """Question box; a bold line right under a question (options, extra text) joins the same box."""
    if continues:
        doc.paragraphs[-1].paragraph_format.space_after = Pt(0)
    paragraph = doc.add_paragraph(style=ds.QUESTION_STYLE if kind == "question" else ds.SUBQUESTION_STYLE)
    if continues:
        paragraph.paragraph_format.space_before = Pt(0)
    points = POINTS.search(text)
    add_inline(paragraph, text[:points.start()] if points else text, report)
    if points:  # "(2 pts)" in small grey, never split over two lines
        paragraph.add_run(NBSP)
        ds.style_runs(add_inline(paragraph, f"({points.group(1).replace(' ', NBSP)})", report),
                      size=max(settings.size - 1.5, 8), color=theme.muted, bold=False)


def add_table(doc, rows: list[list[str]], settings: Settings, theme: ds.Theme, report: Report) -> None:
    width = max(len(row) for row in rows)
    rows = [row + [""] * (width - len(row)) for row in rows]
    table = doc.add_table(rows=len(rows), cols=width)
    for row_index, row in enumerate(rows):
        for col_index, text in enumerate(row):
            add_inline(table.cell(row_index, col_index).paragraphs[0], text, report)
    body = rows[1:]
    centered = {col for col in range(width)
                if any(row[col] for row in body) and all(NUMERIC_CELL.match(row[col]) for row in body if row[col])}
    ds.style_table(table, ds.column_widths(rows, ds.text_width_cm(settings.lang)), theme, centered)
    ds.exact_spacer(doc, 10)


def resolve_image(path: Path, base: Path) -> Path | None:
    """Find an image whether its path is relative to the Markdown file, the current folder or their parent."""
    if path.is_absolute():
        return path if path.is_file() else None
    for root in (base, Path.cwd(), base.parent):
        candidate = (root / path).resolve()
        if candidate.is_file():
            return candidate
    return None


def prepare_image(source: Path, temp_dir: Path, number: int) -> Path | None:
    """Copy the image under a neutral name (no dashes end up in the file), turn phone photos upright
    (EXIF orientation) and convert formats Word can't read. None if it can't be used at all."""
    suffix = source.suffix.lower()
    copied = temp_dir / f"figure{number}{suffix}"
    try:
        from PIL import Image, ImageOps
        with Image.open(source) as image:
            turned = image.getexif().get(0x0112, 1) not in (0, 1)
            if suffix in WORD_IMAGE_TYPES and not turned:
                shutil.copyfile(source, copied)
                return copied
            upright = ImageOps.exif_transpose(image)
            if suffix in (".jpg", ".jpeg"):
                upright.convert("RGB").save(copied, "JPEG", quality=92)
                return copied
            target = temp_dir / f"figure{number}.png"
            upright.convert("RGBA" if upright.mode in ("RGBA", "LA", "P") else "RGB").save(target, "PNG")
            return target
    except Exception:  # no Pillow, or a file Pillow can't open (svg, heic, damaged): let Word try the original
        if suffix in WORD_IMAGE_TYPES:
            shutil.copyfile(source, copied)
            return copied
        return None


def add_caption(doc, caption: str, report: Report) -> None:
    paragraph = doc.add_paragraph(style="Caption")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline(paragraph, caption, report)


def add_image(doc, caption: str, path_text: str, width_cm: float | None, base: Path,
              temp_dir: Path, settings: Settings, report: Report) -> bool:
    source = resolve_image(Path(path_text), base)
    if source is None:
        report.image_problems.append(f"not found: {path_text} (paths are relative to the Markdown file)")
        return False
    prepared = prepare_image(source, temp_dir, len(doc.inline_shapes) + 1)
    paragraphs_before = len(doc.paragraphs)
    try:
        if prepared is None:
            raise ValueError("unsupported format")
        doc.add_picture(str(prepared))
    except Exception as error:
        if len(doc.paragraphs) > paragraphs_before:  # add_picture leaves its empty paragraph behind
            empty = doc.paragraphs[-1]._p
            empty.getparent().remove(empty)
        reason = str(error) or type(error).__name__
        report.image_problems.append(f"unreadable: {path_text} ({reason}); save it as PNG or JPG")
        return False
    picture = doc.inline_shapes[-1]
    natural_width, natural_height = picture.width, picture.height
    width = Cm(min(width_cm, MAX_IMAGE_WIDTH_CM)) if width_cm else min(natural_width, Cm(DEFAULT_IMAGE_WIDTH_CM))
    height = natural_height * width / natural_width
    max_height = Cm(min(MAX_IMAGE_HEIGHT_CM, ds.usable_height_pt(settings.lang) / ds.PT_PER_CM - 5))
    if height > max_height:  # tall phone screenshots and portrait photos: fit the page, keep the shape
        width, height = width * max_height / height, max_height
    picture.width, picture.height = int(width), int(height)
    holder = doc.paragraphs[-1]
    holder.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ds.single_spaced(holder, before=6, after=0 if caption else 10)
    holder.paragraph_format.keep_with_next = bool(caption)
    if caption:
        add_caption(doc, caption, report)
    return True


def add_picture_placeholder(doc, caption: str, hint: str, width_cm: float | None, height_cm: float | None,
                            theme: ds.Theme, report: Report) -> None:
    width = min(width_cm or PLACEHOLDER_SIZE_CM[0], MAX_IMAGE_WIDTH_CM)
    height = min(max(height_cm or PLACEHOLDER_SIZE_CM[1], 3.0), 20.0)
    ds.exact_spacer(doc, 6)
    ds.picture_placeholder(doc, hint, width, height, theme, lambda p, text: add_inline(p, text, report))
    report.pictures.append(clean_heading(strip_dashes(caption or hint or "picture without caption", report)))
    if caption:
        add_caption(doc, caption, report)
    else:
        ds.exact_spacer(doc, 10)


def size_attribute(attrs: str | None, name: str) -> float | None:
    match = re.search(rf"{name}\s*=\s*(\d+(?:\.\d+)?)", attrs or "")
    return float(match.group(1)) if match else None


def add_figure(doc, image: re.Match, base: Path, temp_dir: Path, settings: Settings, theme: ds.Theme,
               report: Report) -> None:
    """A real image, or a placeholder box when the student has to add the picture (or the file is unusable)."""
    caption, path_text, hint = image.group("caption"), image.group("path").strip(), image.group("hint") or ""
    width, height = size_attribute(image.group("attrs"), "width"), size_attribute(image.group("attrs"), "height")
    if path_text.lower() in PLACEHOLDER_PATHS:
        add_picture_placeholder(doc, caption, hint, width, height, theme, report)
    elif not add_image(doc, caption, path_text, width, base, temp_dir, settings, report):
        add_picture_placeholder(doc, caption, hint, width, height, theme, report)


def add_reference(doc, text: str, report: Report) -> None:
    paragraph = doc.add_paragraph(style=ds.REFERENCE_STYLE)
    paragraph.paragraph_format.left_indent = Cm(HANGING_INDENT_CM)
    paragraph.paragraph_format.first_line_indent = Cm(-HANGING_INDENT_CM)
    add_inline(paragraph, re.sub(r"^([-*+•]|\d{1,3}[.)])\s+", "", text), report)


def add_numbered(doc, raw: str, label: str, text: str, theme: ds.Theme, report: Report) -> None:
    indent = Cm(0.75 + (0.75 if raw.startswith("  ") else 0))
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = indent
    paragraph.paragraph_format.first_line_indent = Cm(-0.75)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.tab_stops.add_tab_stop(indent)
    ds.style_runs(add_inline(paragraph, label, report), color=theme.accent, bold=True)
    paragraph.add_run("\t")
    add_inline(paragraph, text, report)


def read_table(lines: list[str], index: int, first: str) -> tuple[list[list[str]], int]:
    rows = [split_row(first)]
    while index < len(lines) and lines[index].strip().startswith("|"):
        if not TABLE_SEPARATOR.match(lines[index].strip()):
            rows.append(split_row(lines[index]))
        index += 1
    return rows, index


def render_body(doc, lines: list[str], base: Path, temp_dir: Path, settings: Settings, theme: ds.Theme,
                report: Report) -> None:
    index, references_level, previous = 0, 0, ""  # references_level > 0 inside a references section
    while index < len(lines):
        raw = lines[index].rstrip().expandtabs(4)
        stripped = raw.strip()
        index += 1
        if not stripped or stripped in {"---", "***"}:  # a lone ___ is a fill-in blank, so it stays
            continue
        kind = "other"
        heading, image, loose_image = HEADING.match(stripped), IMAGE.match(stripped), LOOSE_IMAGE.match(stripped)
        bold_line = FULL_BOLD.match(stripped)
        bullet, numbered = re.match(r"^([-*+•])\s+(.*)$", stripped), NUMBERED.match(stripped)
        question = question_kind(bold_line.group(1), previous) if bold_line and not references_level else None
        if stripped.lower() in PAGEBREAKS:
            page_break(doc)
        elif heading:
            level = len(heading.group(1))
            add_inline(doc.add_paragraph(style=f"Heading {min(level, 3)}"), heading.group(2), report)
            if REFERENCES.match(clean_heading(heading.group(2))):
                references_level = level
            elif references_level and level <= references_level:
                references_level = 0
        elif stripped.startswith("|"):
            rows, index = read_table(lines, index, stripped)
            add_table(doc, rows, settings, theme, report)
        elif image:
            add_figure(doc, image, base, temp_dir, settings, theme, report)
        elif loose_image:
            report.image_problems.append(f"could not read the image line {stripped[:60]!r}")
            add_picture_placeholder(doc, loose_image.group("caption"), "", None, None, theme, report)
        elif references_level:
            add_reference(doc, stripped, report)
        elif question and bold_line:
            continues = question == "question" and previous == "question" and not QUESTION_START.match(
                bold_line.group(1))
            add_question(doc, bold_line.group(1), question, continues, theme, settings, report)
            kind = question
        elif bullet:
            nested = len(raw) - len(raw.lstrip()) >= 2
            add_inline(doc.add_paragraph(style="List Bullet 2" if nested else "List Bullet"), bullet.group(2), report)
        elif numbered:
            add_numbered(doc, raw, numbered.group(1), numbered.group(2), theme, report)
        elif stripped.startswith(">"):
            add_inline(doc.add_paragraph(style=ds.CALLOUT_STYLE), stripped.lstrip("> ").strip(), report)
        else:
            add_inline(doc.add_paragraph(), stripped, report)
        previous = kind


# ------------------------------------------------------------------ assembly

def compose(settings: Settings, details: list[tuple[str, str]], body: list[str], base: Path,
            toc_pages: dict[int, int] | None) -> tuple[WordDocument, Report]:
    report = Report()
    theme = ds.THEMES[settings.theme]
    use_cover = bool(settings.title) and (settings.cover if settings.cover is not None else True)
    details = [(clean_heading(strip_dashes(key, report)), value) for key, value in details]
    doc = Document()
    ds.modern_layout(doc)
    ds.setup_styles(doc, theme, settings.font, settings.size, settings.spacing, settings.lang)
    if use_cover:
        add_cover(doc, settings, details, theme, report)
    else:
        add_header_block(doc, settings, details, theme, report)
    if settings.toc:
        add_toc(doc, settings, collect_headings(body), toc_pages, report)
    with tempfile.TemporaryDirectory() as temp:
        render_body(doc, body, base, Path(temp), settings, theme, report)
    add_page_furniture(doc, settings, details, theme, use_cover)
    return doc, report


def toc_page_numbers(output: Path, body: list[str]) -> tuple[dict[int, int], int] | None:
    """(page of each contents entry, total pages). Every heading level is matched in order, so a ###
    heading with the same text as a later ## heading can't steal its page number."""
    every = collect_headings(body, max_level=6)
    rendered = office_render.heading_pages(output.resolve(), [text for _, text in every])
    if not rendered:
        return None
    found, total_pages = rendered
    pages, entry = {}, 0
    for index, (level, _) in enumerate(every):
        if level <= 2:
            if index in found:
                pages[entry] = found[index]
            entry += 1
    return (pages, total_pages) if pages else None


def body_texts(doc: WordDocument) -> list[str]:
    """Text of every paragraph in the main body, tables included (what Word counts in File > Info)."""
    return ["".join(node.text or "" for node in p.iter(qn("w:t"))) for p in doc.element.body.iter(qn("w:p"))]


def customized_styles(doc: WordDocument) -> frozenset[str]:
    """Styles this build added or changed. Word keeps those even before anything uses them."""
    def snapshot(document: WordDocument) -> dict[str, bytes]:
        return {style.get(qn("w:styleId")): etree.tostring(style)
                for style in document.styles.element.iter(qn("w:style"))}
    template = snapshot(Document())
    return frozenset(sid for sid, xml in snapshot(doc).items() if template.get(sid) != xml)


def finish(doc: WordDocument, output: Path, settings: Settings, details: list[tuple[str, str]],
           use_cover: bool, pages: int | None) -> None:
    doc.save(str(output))
    extra_pages = int(use_cover) + int(settings.toc)
    stats = docx_finalize.document_stats(body_texts(doc), pages, extra_pages)
    author = plain_text(detail_value(details, NAME_KEYS))[:METADATA_MAX]
    docx_finalize.finalize(output, author, stats, keep_styles=customized_styles(doc))


def build(source: Path, output: Path, args: argparse.Namespace) -> Report:
    lines = CONTROL_CHARS.sub("", source.read_text(encoding="utf-8-sig")).splitlines()
    raw_settings, details, body = parse_front_matter(lines)
    settings = build_settings(raw_settings, args)
    output.parent.mkdir(parents=True, exist_ok=True)

    use_cover = bool(settings.title) and (settings.cover if settings.cover is not None else True)
    doc, report = compose(settings, details, body, source.parent, None)
    finish(doc, output, settings, details, use_cover, None)
    if settings.toc:
        rendered = None if args.no_render else toc_page_numbers(output, body)
        if rendered:
            pages, total_pages = rendered
            doc, report = compose(settings, details, body, source.parent, pages)
            finish(doc, output, settings, details, use_cover, total_pages)
            report.notes.append(f"contents page numbers filled in from a LibreOffice render ({len(pages)} headings)")
        else:
            report.notes.append("contents page has no page numbers yet: tell the student to right-click it in "
                                "Word and pick Update Field")
    elif not args.no_render:
        rendered_pages = office_render.page_count(output.resolve())  # for the count in the file properties
        if rendered_pages:
            finish(doc, output, settings, details, use_cover, rendered_pages)
    if raw_settings.get("lang") and lang_code(raw_settings) is None:
        report.notes.append(f"lang {raw_settings['lang']!r} not recognised, used en-US (Letter); use en-US or en-GB")
    return report


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("source", type=Path, help="Markdown file with the content")
    parser.add_argument("-o", "--output", type=Path, required=True, help="where to write the .docx")
    parser.add_argument("--theme", help=f"color theme: {', '.join(ds.THEMES)} (overrides front matter)")
    parser.add_argument("--font", help="font name, e.g. Arial (overrides front matter)")
    parser.add_argument("--size", help="font size, e.g. 12 or 12pt")
    parser.add_argument("--spacing", help="line spacing, e.g. 1.5 or 2")
    parser.add_argument("--no-render", action="store_true", help="skip the LibreOffice pass for page numbers")
    args = parser.parse_args()

    if not args.source.is_file():
        sys.exit(f"File not found: {args.source}")
    report = build(args.source, args.output, args)

    print(f"Saved: {args.output}")
    for note in report.notes:
        print(f"NOTE: {note}")
    if report.dashes_fixed:
        print(f"WARNING: removed {report.dashes_fixed} em/en dash(es) automatically. "
              "Rewrite those sentences in the source so they read naturally, then rebuild.")
    for problem in report.image_problems:
        print(f"ERROR: image {problem}. A picture placeholder box was put in its place; fix the path if "
              "the real figure belongs there.")
    if report.pictures:
        print(f"NOTE: {len(report.pictures)} picture placeholder(s) for the student to fill in by hand: "
              + "; ".join(report.pictures))


if __name__ == "__main__":
    main()
