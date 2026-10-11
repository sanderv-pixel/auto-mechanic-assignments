"""Themes, styles and low-level Word XML helpers used by build_docx.py.

Everything visual lives here: colors, heading look, question boxes, callouts, tables, the cover
page, the worksheet header block and the running header and footer. build_docx.py decides *what*
goes into the document; this module decides how it looks.
"""

from __future__ import annotations

from dataclasses import dataclass
from typing import Callable

from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

DETAILS_STYLE = "Student Details"  # name/class lines; check_text.py leaves them out of the word count
TITLE_STYLE = "Assignment Title"  # title of the worksheet header block; also left out of the count
KICKER_STYLE = "Kicker"  # small capitals label above a title; also left out of the count
QUESTION_STYLE = "Question"
SUBQUESTION_STYLE = "Sub Question"
CALLOUT_STYLE = "Callout"
REFERENCE_STYLE = "Reference"
PLACEHOLDER_STYLE = "Picture Placeholder"  # text inside a picture box; not counted as words

PT_PER_CM = 28.3465
TWIPS_PER_CM = 567
PAGE_SIZES_CM = {"en-US": (21.59, 27.94, 2.54), "en-GB": (21.0, 29.7, 2.5)}  # width, height, margin


@dataclass(frozen=True)
class Theme:
    primary: str  # title, H1, H3, question text, table header
    accent: str  # subtitle, H2, rules, kicker, step numbers
    tint: str  # question box and info bar fill
    band: str  # alternate table rows, callout fill
    line: str  # table borders, header rule
    text: str
    muted: str  # labels, captions, header and footer
    header_fill: str
    header_text: str


THEMES = {
    "navy": Theme(primary="1F3864", accent="2E75B6", tint="EAF1FB", band="F3F6FB", line="C5D0E0",
                  text="1E1E1E", muted="6B7280", header_fill="1F3864", header_text="FFFFFF"),
    "graphite": Theme(primary="2B2B2B", accent="B3261E", tint="F3F3F3", band="F7F7F7", line="D0D0D0",
                      text="1E1E1E", muted="6B6B6B", header_fill="2B2B2B", header_text="FFFFFF"),
    "plain": Theme(primary="000000", accent="000000", tint="FFFFFF", band="FFFFFF", line="7F7F7F",
                   text="000000", muted="404040", header_fill="E6E6E6", header_text="000000"),
}

PPR_ORDER = ["w:pStyle", "w:keepNext", "w:keepLines", "w:pageBreakBefore", "w:framePr", "w:widowControl",
             "w:numPr", "w:suppressLineNumbers", "w:pBdr", "w:shd", "w:tabs", "w:suppressAutoHyphens",
             "w:kinsoku", "w:wordWrap", "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN",
             "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing",
             "w:mirrorIndents", "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment",
             "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange"]
RPR_ORDER = ["w:rStyle", "w:rFonts", "w:b", "w:bCs", "w:i", "w:iCs", "w:caps", "w:smallCaps", "w:strike",
             "w:dstrike", "w:outline", "w:shadow", "w:emboss", "w:imprint", "w:noProof", "w:snapToGrid",
             "w:vanish", "w:webHidden", "w:color", "w:spacing", "w:w", "w:kern", "w:position", "w:sz", "w:szCs",
             "w:highlight", "w:u", "w:effect", "w:bdr", "w:shd", "w:fitText", "w:vertAlign", "w:rtl", "w:cs",
             "w:em", "w:lang", "w:eastAsianLayout", "w:specVanish", "w:oMath"]
TCPR_ORDER = ["w:cnfStyle", "w:tcW", "w:gridSpan", "w:hMerge", "w:vMerge", "w:tcBorders", "w:shd", "w:noWrap",
              "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark"]
TBLPR_ORDER = ["w:tblStyle", "w:tblpPr", "w:tblOverlap", "w:bidiVisual", "w:tblStyleRowBandSize",
               "w:tblStyleColBandSize", "w:tblW", "w:jc", "w:tblCellSpacing", "w:tblInd", "w:tblBorders", "w:shd",
               "w:tblLayout", "w:tblCellMar", "w:tblLook"]
TRPR_ORDER = ["w:cnfStyle", "w:divId", "w:gridBefore", "w:gridAfter", "w:wBefore", "w:wAfter", "w:cantSplit",
              "w:trHeight", "w:tblHeader", "w:tblCellSpacing", "w:jc", "w:hidden"]

AddText = Callable[[object, str], list]  # (paragraph, markdown text) -> runs


# ---------------------------------------------------------------- XML helpers

def set_child(parent, tag: str, order: list[str]):
    """Replace (or create) a child element, inserted where the schema wants it."""
    existing = parent.find(qn(tag))
    if existing is not None:
        parent.remove(existing)
    element = OxmlElement(tag)
    successors = {qn(name) for name in order[order.index(tag) + 1:]}
    for child in parent:
        if child.tag in successors:
            child.addprevious(element)
            return element
    parent.append(element)
    return element


def ppr_of(target):
    """pPr of a paragraph or of a paragraph style."""
    return target._p.get_or_add_pPr() if hasattr(target, "_p") else target.element.get_or_add_pPr()


def set_borders(target, **sides: tuple[int, str, int]) -> None:
    """sides: top/left/bottom/right=(width in eighths of a point, hex color, gap to the text in pt)."""
    borders = set_child(ppr_of(target), "w:pBdr", PPR_ORDER)
    for side in ("top", "left", "bottom", "right"):
        if side in sides:
            size, color, space = sides[side]
            edge = OxmlElement(f"w:{side}")
            edge.set(qn("w:val"), "single")
            edge.set(qn("w:sz"), str(size))
            edge.set(qn("w:space"), str(space))
            edge.set(qn("w:color"), color)
            borders.append(edge)


def set_shading(target, fill: str) -> None:
    shd = set_child(ppr_of(target), "w:shd", PPR_ORDER)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)


def set_letter_spacing(run_or_style, twentieths_of_point: int) -> None:
    owner = run_or_style._r if hasattr(run_or_style, "_r") else run_or_style.element
    spacing = set_child(owner.get_or_add_rPr(), "w:spacing", RPR_ORDER)
    spacing.set(qn("w:val"), str(twentieths_of_point))


def set_language(rpr, lang: str) -> None:
    lang_el = set_child(rpr, "w:lang", RPR_ORDER)
    lang_el.set(qn("w:val"), lang)
    lang_el.set(qn("w:eastAsia"), lang)


def set_font(style, font: str) -> None:
    style.font.name = font
    rfonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        rfonts.attrib.pop(qn(attr), None)  # theme fonts would win over the font we set (headings, Title)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), font)


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def style_runs(runs: list, size: float | None = None, color: str | None = None,
               bold: bool | None = None, italic: bool | None = None) -> None:
    for run in runs:
        if size:
            run.font.size = Pt(size)
        if color:
            run.font.color.rgb = rgb(color)
        if bold is not None:
            run.font.bold = bold
        if italic is not None:
            run.font.italic = italic


def single_spaced(paragraph, before: float = 0, after: float = 0) -> None:
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)


def exact_spacer(container, points: float):
    """An empty paragraph of a fixed height that ignores the body line spacing."""
    paragraph = container.add_paragraph()
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    paragraph.paragraph_format.line_spacing = Pt(max(points, 1))
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    return paragraph


# ---------------------------------------------------------------- fields

def field_start(paragraph, instruction: str) -> None:
    begin, instr, separate = OxmlElement("w:fldChar"), OxmlElement("w:instrText"), OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {instruction} "
    separate.set(qn("w:fldCharType"), "separate")
    for element in (begin, instr, separate):
        paragraph.add_run()._r.append(element)


def field_end(paragraph) -> None:
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    paragraph.add_run()._r.append(end)


def add_field(paragraph, instruction: str, shown: str, size: float, color: str) -> None:
    field_start(paragraph, instruction)
    style_runs([paragraph.add_run(shown)], size=size, color=color)
    field_end(paragraph)


# ---------------------------------------------------------------- page and styles

def text_width_cm(lang: str) -> float:
    width, _, margin = PAGE_SIZES_CM[lang]
    return width - 2 * margin


def usable_height_pt(lang: str) -> float:
    _, height, margin = PAGE_SIZES_CM[lang]
    return (height - 2 * margin) * PT_PER_CM


def modern_layout(doc) -> None:
    """Word 2013+ layout mode, so tables line up with the margins the way LibreOffice draws them."""
    for setting in doc.settings.element.iter(qn("w:compatSetting")):
        if setting.get(qn("w:name")) == "compatibilityMode":
            setting.set(qn("w:val"), "15")


def setup_page(doc, lang: str) -> None:
    width, height, margin = PAGE_SIZES_CM[lang]
    for section in doc.sections:
        section.page_width, section.page_height = Cm(width), Cm(height)
        section.top_margin = section.bottom_margin = Cm(margin)
        section.left_margin = section.right_margin = Cm(margin)
        section.header_distance = section.footer_distance = Cm(1.25)


def paragraph_style(doc, name: str, base: str = "Normal"):
    if name not in [s.name for s in doc.styles]:
        style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = doc.styles[base]
        style.quick_style = True
    return doc.styles[name]


def setup_base(doc, theme: Theme, font: str, size: float, spacing: float, lang: str) -> None:
    normal = doc.styles["Normal"]
    set_font(normal, font)
    normal.font.size = Pt(size)
    normal.font.color.rgb = rgb(theme.text)
    normal.paragraph_format.line_spacing = spacing
    normal.paragraph_format.space_after = Pt(8)
    set_language(normal.element.get_or_add_rPr(), lang)
    defaults = doc.styles.element.find(qn("w:docDefaults"))
    default_rpr = defaults.find(qn("w:rPrDefault")) if defaults is not None else None
    if default_rpr is not None and default_rpr.find(qn("w:rPr")) is not None:
        set_language(default_rpr.find(qn("w:rPr")), lang)
    for name in ("Title", "Subtitle", "List Bullet", "List Bullet 2", "Caption", "TOC Heading"):
        set_font(doc.styles[name], font)
    for name in ("List Bullet", "List Bullet 2"):
        doc.styles[name].paragraph_format.space_after = Pt(3)


def setup_headings(doc, theme: Theme, font: str, size: float, spacing: float) -> None:
    look = {1: (size + 5, theme.primary, 20, 8), 2: (size + 2, theme.accent, 14, 4),
            3: (size + 0.5, theme.primary, 10, 3)}
    for level, (points, color, before, after) in look.items():
        style = doc.styles[f"Heading {level}"]
        set_font(style, font)
        style.font.size = Pt(points)
        style.font.bold, style.font.italic = True, False
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.0 if spacing < 2 else spacing
        style.paragraph_format.keep_with_next = True
    set_borders(doc.styles["Heading 1"], bottom=(6, theme.accent, 4))

    title = doc.styles["Title"]
    set_borders(title)  # the default template draws a rule under Title; the cover adds its own
    title.font.size, title.font.bold, title.font.color.rgb = Pt(32), True, rgb(theme.primary)
    subtitle = doc.styles["Subtitle"]
    subtitle.font.size, subtitle.font.italic, subtitle.font.color.rgb = Pt(15), False, rgb(theme.accent)

    header_title = paragraph_style(doc, TITLE_STYLE)
    header_title.font.size, header_title.font.bold = Pt(22), True
    header_title.font.color.rgb = rgb(theme.primary)
    header_title.paragraph_format.line_spacing = 1.0
    header_title.paragraph_format.space_after = Pt(4)

    kicker = paragraph_style(doc, KICKER_STYLE)
    kicker.font.size, kicker.font.bold, kicker.font.all_caps = Pt(9.5), True, True
    kicker.font.color.rgb = rgb(theme.accent)
    kicker.paragraph_format.line_spacing = 1.0
    kicker.paragraph_format.space_after = Pt(6)
    set_letter_spacing(kicker, 30)

    details = paragraph_style(doc, DETAILS_STYLE)
    details.paragraph_format.line_spacing = 1.0
    details.paragraph_format.space_after = Pt(0)


def setup_toc_styles(doc, theme: Theme, size: float, lang: str) -> None:
    toc_heading = doc.styles["TOC Heading"]
    toc_heading.font.size, toc_heading.font.bold = Pt(size + 5), True
    toc_heading.font.color.rgb = rgb(theme.primary)
    toc_heading.paragraph_format.space_before = Pt(0)
    toc_heading.paragraph_format.space_after = Pt(14)
    set_borders(toc_heading, bottom=(6, theme.accent, 4))
    for level in (1, 2):
        style = paragraph_style(doc, f"toc {level}")
        style.font.bold = level == 1
        style.font.color.rgb = rgb(theme.primary if level == 1 else theme.text)
        style.paragraph_format.line_spacing = 1.0
        style.paragraph_format.space_before = Pt(8 if level == 1 else 0)
        style.paragraph_format.space_after = Pt(3)
        style.paragraph_format.left_indent = Cm(0 if level == 1 else 0.6)
        style.paragraph_format.tab_stops.add_tab_stop(Cm(text_width_cm(lang)), WD_TAB_ALIGNMENT.RIGHT,
                                                      WD_TAB_LEADER.DOTS)


def setup_blocks(doc, theme: Theme, size: float) -> None:
    question = paragraph_style(doc, QUESTION_STYLE)
    question.font.bold = True
    question.font.color.rgb = rgb(theme.primary)
    fmt = question.paragraph_format
    fmt.line_spacing, fmt.space_before, fmt.space_after = 1.1, Pt(16), Pt(6)
    fmt.keep_with_next, fmt.left_indent = True, Cm(0.35)
    set_shading(question, theme.tint)
    set_borders(question, top=(2, theme.tint, 4), left=(24, theme.primary, 7), bottom=(2, theme.tint, 4))

    sub_question = paragraph_style(doc, SUBQUESTION_STYLE)
    sub_question.font.bold = True
    sub_question.font.color.rgb = rgb(theme.primary)
    fmt = sub_question.paragraph_format
    fmt.line_spacing, fmt.space_before, fmt.space_after = 1.1, Pt(8), Pt(4)
    fmt.keep_with_next, fmt.left_indent = True, Cm(0.35)
    set_borders(sub_question, left=(12, theme.accent, 7))

    callout = paragraph_style(doc, CALLOUT_STYLE)
    fmt = callout.paragraph_format
    fmt.left_indent = Cm(0.4)
    fmt.space_before, fmt.space_after = Pt(4), Pt(10)
    set_shading(callout, theme.band)
    set_borders(callout, top=(2, theme.band, 5), left=(18, theme.accent, 8), bottom=(2, theme.band, 5))

    reference = paragraph_style(doc, REFERENCE_STYLE)
    reference.font.size = Pt(max(size - 1, 9))
    reference.paragraph_format.space_after = Pt(6)

    placeholder = paragraph_style(doc, PLACEHOLDER_STYLE)
    placeholder.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    placeholder.paragraph_format.line_spacing = 1.0
    placeholder.paragraph_format.space_after = Pt(4)
    placeholder.paragraph_format.keep_with_next = True

    caption = doc.styles["Caption"]
    caption.font.size, caption.font.italic, caption.font.bold = Pt(9), True, False
    caption.font.color.rgb = rgb(theme.muted)
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(12)


def setup_styles(doc, theme: Theme, font: str, size: float, spacing: float, lang: str) -> None:
    setup_page(doc, lang)
    setup_base(doc, theme, font, size, spacing, lang)
    setup_headings(doc, theme, font, size, spacing)
    setup_toc_styles(doc, theme, size, lang)
    setup_blocks(doc, theme, size)


# ---------------------------------------------------------------- tables

def table_width(table, total_cm: float) -> None:
    width = set_child(table._tbl.tblPr, "w:tblW", TBLPR_ORDER)
    width.set(qn("w:w"), str(round(total_cm * TWIPS_PER_CM)))
    width.set(qn("w:type"), "dxa")


def table_borders(table, color: str) -> None:
    borders = set_child(table._tbl.tblPr, "w:tblBorders", TBLPR_ORDER)
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = OxmlElement(f"w:{side}")
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "4")
        edge.set(qn("w:space"), "0")
        edge.set(qn("w:color"), color)
        borders.append(edge)


def table_cell_margins(table, top_bottom_twips: int, left_right_twips: int) -> None:
    margins = set_child(table._tbl.tblPr, "w:tblCellMar", TBLPR_ORDER)
    for side, value in (("top", top_bottom_twips), ("left", left_right_twips),
                        ("bottom", top_bottom_twips), ("right", left_right_twips)):
        edge = OxmlElement(f"w:{side}")
        edge.set(qn("w:w"), str(value))
        edge.set(qn("w:type"), "dxa")
        margins.append(edge)


def cell_shading(cell, fill: str) -> None:
    shd = set_child(cell._tc.get_or_add_tcPr(), "w:shd", TCPR_ORDER)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)


def cell_borders(cell, sides: dict[str, tuple[int, str]], line: str = "single") -> None:
    """Borders for one cell, sides = {"top": (width in eighths of a point, color), ...}; others switched off."""
    borders = set_child(cell._tc.get_or_add_tcPr(), "w:tcBorders", TCPR_ORDER)
    for side in ("top", "left", "bottom", "right"):
        edge = OxmlElement(f"w:{side}")
        if side in sides:
            size, color = sides[side]
            edge.set(qn("w:val"), line)
            edge.set(qn("w:sz"), str(size))
            edge.set(qn("w:space"), "0")
            edge.set(qn("w:color"), color)
        else:
            edge.set(qn("w:val"), "nil")
        borders.append(edge)


def repeat_header_row(row) -> None:
    header = set_child(row._tr.get_or_add_trPr(), "w:tblHeader", TRPR_ORDER)
    header.set(qn("w:val"), "true")


def fixed_grid(table, widths_cm: list[float]) -> None:
    table.autofit = False
    table_width(table, sum(widths_cm))
    for index, width in enumerate(widths_cm):
        table.columns[index].width = Cm(width)
        for cell in table.columns[index].cells:
            cell.width = Cm(width)


def column_widths(rows: list[list[str]], total_cm: float) -> list[float]:
    """Share the width by content length, so short columns stay narrow and long ones get room."""
    count = max(len(row) for row in rows)
    lengths = [min(max(max(len(row[i]) for row in rows if i < len(row)), 6), 60) for i in range(count)]
    widths = [max(total_cm * length / sum(lengths), total_cm * 0.12) for length in lengths]
    scale = total_cm / sum(widths)
    return [width * scale for width in widths]


def style_table(table, widths_cm: list[float], theme: Theme, centered_columns: set[int]) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    fixed_grid(table, widths_cm)
    table_borders(table, theme.line)
    table_cell_margins(table, 70, 115)
    for row_index, row in enumerate(table.rows):
        is_header = row_index == 0
        if is_header:
            repeat_header_row(row)
        for col_index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if is_header:
                cell_shading(cell, theme.header_fill)
            elif row_index % 2 == 0:
                cell_shading(cell, theme.band)
            for paragraph in cell.paragraphs:
                single_spaced(paragraph)
                if col_index in centered_columns:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                if is_header:
                    style_runs(paragraph.runs, size=10, color=theme.header_text, bold=True)
                else:
                    style_runs(paragraph.runs, size=10)


# ---------------------------------------------------------------- cover, header block, page furniture

def label_run(paragraph, text: str, theme: Theme, size: float = 8) -> None:
    run = paragraph.add_run(text.upper())
    style_runs([run], size=size, color=theme.muted, bold=True)
    set_letter_spacing(run, 20)


def details_table(doc, details: list[tuple[str, str]], theme: Theme, total_cm: float, add_text: AddText) -> None:
    """Cover page details: small grey labels on the left, values on the right, a thin rule per row."""
    table = doc.add_table(rows=len(details), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table_cell_margins(table, 95, 0)
    widths = [4.0, total_cm - 4.0]
    for row_index, (key, value) in enumerate(details):
        cells = table.rows[row_index].cells
        for cell in cells:
            top = {"top": (12, theme.primary)} if row_index == 0 else {}
            cell_borders(cell, {"bottom": (4, theme.line), **top})
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.paragraphs[0].style = doc.styles[DETAILS_STYLE]
        label_run(cells[0].paragraphs[0], key, theme, 8.5)
        add_text(cells[1].paragraphs[0], value)
    fixed_grid(table, widths)


def info_bar(doc, details: list[tuple[str, str]], theme: Theme, total_cm: float, add_text: AddText) -> None:
    """Worksheet header: the details side by side in a tinted band, label above value."""
    per_row = len(details) if len(details) <= 3 else (2 if len(details) == 4 else 3)
    groups = [details[i:i + per_row] for i in range(0, len(details), per_row)]
    if len(groups) == 1:
        weights = [max(len(key), len(value), 10) for key, value in details]
        widths = [total_cm * weight / sum(weights) for weight in weights]
    else:
        widths = [total_cm / per_row] * per_row
    table = doc.add_table(rows=len(groups), cols=per_row)
    table_cell_margins(table, 85, 140)
    for row_index, group in enumerate(groups):
        for col_index, cell in enumerate(table.rows[row_index].cells):
            cell_shading(cell, theme.tint)
            cell_borders(cell, {"left": (24, theme.primary)} if col_index == 0 else {})
            first = cell.paragraphs[0]
            first.style = doc.styles[DETAILS_STYLE]
            if col_index >= len(group):
                continue
            key, value = group[col_index]
            label_run(first, key, theme)
            value_paragraph = cell.add_paragraph(style=doc.styles[DETAILS_STYLE])
            value_paragraph.paragraph_format.space_before = Pt(2)
            add_text(value_paragraph, value)
    fixed_grid(table, widths)


def page_header(section, left: str, right: str, theme: Theme, total_cm: float) -> None:
    header = section.header
    header.is_linked_to_previous = False
    paragraph = header.paragraphs[0]
    paragraph.style.paragraph_format.tab_stops.clear_all()  # the template's center tab would catch the name
    paragraph.paragraph_format.tab_stops.add_tab_stop(Cm(total_cm), WD_TAB_ALIGNMENT.RIGHT)
    set_borders(paragraph, bottom=(4, theme.line, 4))
    style_runs([paragraph.add_run(f"{left}\t{right}")], size=8.5, color=theme.muted)


def page_footer(footer, theme: Theme) -> None:
    footer.is_linked_to_previous = False
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_runs([paragraph.add_run("Page ")], size=8.5, color=theme.muted)
    add_field(paragraph, "PAGE", "1", 8.5, theme.muted)
    style_runs([paragraph.add_run(" of ")], size=8.5, color=theme.muted)
    add_field(paragraph, "NUMPAGES", "1", 8.5, theme.muted)
    style_runs(paragraph.runs, size=8.5, color=theme.muted)  # LibreOffice formats a field like its first run


def picture_placeholder(doc, hint: str, width_cm: float, height_cm: float, theme: Theme, add_text: AddText) -> None:
    """A dashed box the student clicks into to paste their own photo or drawing."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    fixed_grid(table, [width_cm])
    row = table.rows[0]
    row.height, row.height_rule = Cm(height_cm), WD_ROW_HEIGHT_RULE.AT_LEAST
    set_child(row._tr.get_or_add_trPr(), "w:cantSplit", TRPR_ORDER)
    cell = row.cells[0]
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    cell_shading(cell, theme.band)
    edge = (12, theme.accent)
    cell_borders(cell, {"top": edge, "left": edge, "bottom": edge, "right": edge}, line="dashed")
    label = cell.paragraphs[0]
    label.style = doc.styles[PLACEHOLDER_STYLE]
    label_run(label, "Insert picture here", theme, 10)
    style_runs(label.runs, color=theme.accent)
    if hint:
        description = cell.add_paragraph(style=doc.styles[PLACEHOLDER_STYLE])
        style_runs(add_text(description, hint), size=10, color=theme.text, italic=True)
    tip = cell.add_paragraph(style=doc.styles[PLACEHOLDER_STYLE])
    style_runs([tip.add_run("Click in this box, delete this text, then use Insert > Pictures")], size=8,
               color=theme.muted)
