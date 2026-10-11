"""Regression tests for the skill's scripts (build_docx, check_text, extract_pdf).

Run: python -m pytest tests/ -q   (needs python-docx, pymupdf, pillow, pytest)
"""

from __future__ import annotations

import subprocess
import sys
import zipfile
from pathlib import Path

import pytest
from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn

SCRIPTS = Path(__file__).resolve().parents[1] / "auto-mechanic-assignments" / "scripts"
sys.path.insert(0, str(SCRIPTS))
import check_text  # noqa: E402
import office_render  # noqa: E402

EM, EN = chr(0x2014), chr(0x2013)
DASH_CHARS = (EM, EN, chr(0x2015), chr(0x2012))


def build(tmp_path: Path, markdown: str, *flags: str) -> tuple[Path, str]:
    source = tmp_path / "answers.md"
    source.write_text(markdown, encoding="utf-8")
    output = tmp_path / "out.docx"
    run = subprocess.run([sys.executable, str(SCRIPTS / "build_docx.py"), str(source), "-o", str(output), *flags],
                         capture_output=True, text=True, cwd=tmp_path)
    assert run.returncode == 0, run.stderr
    return output, run.stdout


def check(path: Path, *flags: str) -> subprocess.CompletedProcess:
    return subprocess.run([sys.executable, str(SCRIPTS / "check_text.py"), str(path), *flags],
                          capture_output=True, text=True)


def all_xml(path: Path) -> str:
    with zipfile.ZipFile(path) as archive:
        return "".join(archive.read(n).decode("utf-8", "ignore") for n in archive.namelist() if n.endswith(".xml"))


# ------------------------------------------------------------------ dashes

def test_dashes_removed_everywhere_including_metadata(tmp_path):
    md = (f"---\ntitle: Brake Systems {EN} Assignment 3\nsubtitle: Part A {EM} Basics\ntoc: true\n"
          f"toc_title: Contents {EM} overview\nName: Jake\n---\n# Intro\nBrake fluid {EM} it absorbs water.\n"
          f"Ranges 10{EN}12 mm and a--b and trailing {EM}\nNote: {EM} x\n")
    output, stdout = build(tmp_path, md)
    xml = all_xml(output)
    assert not any(char in xml for char in DASH_CHARS)
    assert "WARNING: removed" in stdout
    assert "Note:, x" not in xml and "Note: x" in xml
    assert check(output).returncode == 0


def test_image_with_dash_in_file_name_is_renamed(tmp_path):
    from PIL import Image
    Image.new("RGB", (40, 20), "white").save(tmp_path / f"Screenshot {EN} copy.png")
    output, _ = build(tmp_path, f"---\ncover: false\n---\n![Figure 1](Screenshot {EN} copy.png)\n")
    assert len(Document(str(output)).inline_shapes) == 1
    assert EN not in all_xml(output)


def test_checker_finds_dash_hidden_in_properties(tmp_path):
    doc = Document()
    doc.add_paragraph("Clean text.")
    doc.core_properties.title = f"Brakes {EM} test"
    doc.core_properties.author = "Jake"
    path = tmp_path / "hidden.docx"
    doc.save(str(path))
    result = check(path)
    assert result.returncode == 1 and "dash hidden in docProps/core.xml" in result.stdout


def test_double_hyphen_flagged_in_markdown(tmp_path):
    path = tmp_path / "a.md"
    path.write_text("The pads wore out--fast.\n", encoding="utf-8")
    assert check(path).returncode == 1


# ------------------------------------------------------------- robustness

def test_units_in_size_and_spacing_are_understood(tmp_path):
    output, _ = build(tmp_path, "---\ntitle: T\ncover: false\nsize: 12pt\nspacing: 1,5\n---\nText.\n")
    normal = Document(str(output)).styles["Normal"]
    assert normal.font.size.pt == 12 and normal.paragraph_format.line_spacing == 1.5


def test_bad_images_do_not_abort_the_build(tmp_path):
    from PIL import Image
    Image.new("RGB", (40, 20), "red").save(tmp_path / "pic.webp")
    (tmp_path / "broken.png").write_bytes(b"not an image")
    output, stdout = build(tmp_path, "---\ncover: false\n---\n![A](pic.webp)\n![B](broken.png)\n"
                                     "![C](missing.png)\nEnd.\n")
    assert len(Document(str(output)).inline_shapes) == 1  # webp converted, the other two reported
    assert "unreadable: broken.png" in stdout and "not found: missing.png" in stdout


def test_image_width_is_clamped_and_path_relative_to_markdown(tmp_path):
    from PIL import Image
    (tmp_path / "work").mkdir()
    Image.new("RGB", (400, 200), "white").save(tmp_path / "work" / "fig.png")
    source = tmp_path / "work" / "answers.md"
    source.write_text("![Fig](work/fig.png){width=25}\n", encoding="utf-8")
    output = tmp_path / "out.docx"
    subprocess.run([sys.executable, str(SCRIPTS / "build_docx.py"), str(source), "-o", str(output)],
                   check=True, cwd=tmp_path / "work", capture_output=True)
    assert round(Document(str(output)).inline_shapes[0].width.cm) == 16


def test_bom_front_matter_and_capitalized_setting_keys(tmp_path):
    source = tmp_path / "answers.md"
    source.write_text("---\nTitle: Cooling\nCover: false\nName: [Name]\n---\nText.\n", encoding="utf-8-sig")
    output = tmp_path / "out.docx"
    subprocess.run([sys.executable, str(SCRIPTS / "build_docx.py"), str(source), "-o", str(output)], check=True)
    texts = [p.text for p in Document(str(output)).paragraphs]
    assert texts[0] == "Cooling" and "---" not in texts and "Title: Cooling" not in texts


# ---------------------------------------------------------------- markdown

def test_markdown_inline_handling(tmp_path):
    md = ("---\ncover: false\n---\n#### Deep heading\n***Warning*** read `P0301` first.\n"
          "Fill in ___ and ___ please.\nSee https://example.com/ASE_xEV_Safety_Standards.pdf now.\n")
    output, _ = build(tmp_path, md)
    paragraphs = Document(str(output)).paragraphs
    heading = next(p for p in paragraphs if "Deep heading" in p.text)
    assert heading.style.name == "Heading 3" and "#" not in heading.text
    warning = next(p for p in paragraphs if "Warning" in p.text)
    assert warning.text == "Warning read P0301 first."
    assert warning.runs[0].bold and warning.runs[0].italic
    assert any(p.text == "Fill in ___ and ___ please." for p in paragraphs)
    assert any("ASE_xEV_Safety_Standards.pdf" in p.text for p in paragraphs)


def test_cover_keeps_its_own_spacing_at_double_spacing(tmp_path):
    output, _ = build(tmp_path, "---\ntitle: Paper\nsubtitle: Sub\nName: A\nClass: B\nDate: C\n---\n# Intro\nText.\n",
                      "--spacing", "2")
    doc = Document(str(output))
    paragraphs = doc.paragraphs
    cover = paragraphs[:paragraphs.index(next(p for p in paragraphs if p.text == "Intro"))]
    for paragraph in cover:  # text lines single spaced, spacers a fixed height: nothing grows with the setting
        fmt = paragraph.paragraph_format
        spacer = not paragraph.text.strip() and fmt.line_spacing_rule == WD_LINE_SPACING.EXACTLY
        assert fmt.line_spacing == 1.0 or spacer
    assert 'w:type="page"' in cover[-1]._p.xml
    assert [c.text for c in doc.tables[0].rows[0].cells] == ["NAME", "A"]


needs_libreoffice = pytest.mark.skipif(office_render.find_soffice() is None, reason="LibreOffice not installed")


@needs_libreoffice
def test_long_cover_fits_on_one_page_and_toc_gets_page_numbers(tmp_path):
    details = "".join(f"Detail {n}: value number {n}\n" for n in range(8))
    title = "Diagnosing an Overheating Engine Cooling System on a Pickup Truck in the School Workshop"
    front = f"---\ntitle: {title}\nsubtitle: Shop report\nkicker: Report\ntoc: true\n{details}---\n"
    output, stdout = build(tmp_path, front +
                           "# Intro\nText.\n\\pagebreak\n# Results\nMore text.\n", "--spacing", "2")
    assert "page numbers filled in" in stdout
    pdf = office_render.to_pdf(output, tmp_path / "pdf")
    with office_render.load_pymupdf().open(str(pdf)) as rendered:
        assert "value number 7" in rendered[0].get_text() and "Contents" in rendered[1].get_text()
    toc = [p.text for p in Document(str(output)).paragraphs if p.style.name.startswith("toc")]
    assert toc == ["Intro\t3", "Results\t4"]


def test_toc_without_render_has_no_numbers(tmp_path):
    output, stdout = build(tmp_path, "---\ntitle: T\ntoc: true\n---\n# Intro\nText.\n", "--no-render")
    assert "Update Field" in stdout
    assert [p.text for p in Document(str(output)).paragraphs if p.style.name.startswith("toc")] == ["Intro"]


def test_en_gb_uses_a4(tmp_path):
    output, _ = build(tmp_path, "---\nlang: en-GB\ncover: false\n---\nText.\n")
    section = Document(str(output)).sections[0]
    assert round(section.page_width.cm, 1) == 21.0 and round(section.page_height.cm, 1) == 29.7


def test_references_with_sub_headings_keep_hanging_indent(tmp_path):
    md = "---\ncover: false\n---\n# Body\nOne two three.\n# References\n### Books\nHillier, V. (2012). Book.\n"
    output, _ = build(tmp_path, md)
    reference = next(p for p in Document(str(output)).paragraphs if p.text.startswith("Hillier"))
    assert round(reference.paragraph_format.first_line_indent.cm, 2) == -1.27
    assert "words (Word-style, without cover, contents, references): 4" in check(output).stdout  # "Body" + 3 words


# -------------------------------------------------------------- word count

def test_word_count_matches_word_for_measurements():
    assert check_text.count_words("3.5L 12.6V 0.5mm 10°C e.g. 1,000-mile = × ___") == 6


def test_markdown_and_docx_counts_agree(tmp_path):
    from PIL import Image
    Image.new("RGB", (40, 20), "white").save(tmp_path / "none.png")
    md = ("---\ntitle: Assignment 3\nsubtitle: Brakes\ncover: false\nName: [Name]\n---\n"
          "**1. What does ABS do?**\nIt stops the wheels locking.\n1. Raise the car\na) part one\n- bullet item\n"
          "| Part | Job |\n|---|---|\n| Pad | Grips the rotor |\n![Figure 1: Disc brake](none.png)\n"
          "# References\nSmith, J. (2020). Brakes. https://example.com/a_b\n")
    output, _ = build(tmp_path, md)
    md_count = check(tmp_path / "answers.md").stdout.split("): ")[1].split()[0]
    docx_count = check(output).stdout.split("): ")[1].split()[0]
    assert md_count == docx_count


def test_placeholders_listed_but_not_abbreviations(tmp_path):
    output, _ = build(tmp_path, "---\ntitle: T\nName: [Name]\n---\n# A\nPer the AFDC [AFDC] and source [1].\n")
    stdout = check(output).stdout
    assert "[Name]" in stdout and "[AFDC]" not in stdout and "[1]" not in stdout


def test_curly_apostrophe_phrase_is_caught(tmp_path):
    path = tmp_path / "a.md"
    path.write_text("It’s important to note that pads wear.\n", encoding="utf-8")
    assert "it's important to note" in check(path).stdout


# ------------------------------------------------------------------ extract

def test_extract_clears_stale_page_images(tmp_path):
    import fitz
    for name, pages in (("two.pdf", 2), ("one.pdf", 1)):
        doc = fitz.open()
        for number in range(pages):
            doc.new_page().insert_text((72, 72), f"Question {number + 1}: explain ABS in your own words.")
        doc.save(str(tmp_path / name))
    for name in ("two.pdf", "one.pdf"):
        subprocess.run([sys.executable, str(SCRIPTS / "extract_pdf.py"), str(tmp_path / name),
                        "--out", str(tmp_path / "work")], check=True, capture_output=True)
    assert len(list((tmp_path / "work" / "pages").glob("page*.png"))) == 1


TEST_PDFS = ["assignment3_brake_systems.pdf", "research_paper_ev.pdf", "assignment_7_2_cooling.pdf",
             "practical5_brake_pad_job_report.pdf"]


@pytest.mark.parametrize("pdf", TEST_PDFS)
def test_extract_real_test_pdfs(tmp_path, pdf):
    source = Path(__file__).resolve().parents[1] / "test-files" / pdf
    result = subprocess.run([sys.executable, str(SCRIPTS / "extract_pdf.py"), str(source), "--out", str(tmp_path)],
                            capture_output=True, text=True, check=True)
    assert "Backend:" in result.stdout and (tmp_path / "task_text.md").read_text().count("## Page") >= 1


def test_reference_descriptions_are_not_placeholders_and_sections_are_counted(tmp_path):
    path = tmp_path / "a.md"
    path.write_text("# Intro\nOne two three.\n# Part 1\nFour five.\n# References\n"
                    "Tesla. (2024). [Model S service manual]. https://example.com\n", encoding="utf-8")
    stdout = check(path).stdout
    assert "Model S service manual" not in stdout
    assert "words per section: Intro 3 | Part 1 2" in stdout


# ---------------------------------------------------------------- design

def test_question_boxes_sub_questions_and_options(tmp_path):
    md = ("---\ncover: false\n---\n**1. What does ABS do? (2 pts)**\nIt stops lockup.\n"
          "**2. Which fluid is silicone based? (1 pt)**\n**a) DOT 3, b) DOT 5, c) DOT 4**\nb) DOT 5.\n"
          "**3. Calculation (4 pts). A force of 150 lb acts on a piston.**\n**a) Find the pressure.**\nAbout 191 psi.\n"
          "**Bold sentence that is not a question.**\n> **Given:** F = 150 lb\n")
    output, _ = build(tmp_path, md)
    styles = {p.text: p.style.name for p in Document(str(output)).paragraphs if p.text}
    nbsp = chr(0xA0)
    assert styles[f"1. What does ABS do?{nbsp}(2{nbsp}pts)"] == "Question"
    assert styles["a) DOT 3, b) DOT 5, c) DOT 4"] == "Question"  # options join the question box
    assert styles["a) Find the pressure."] == "Sub Question"
    assert styles["Bold sentence that is not a question."] == "Normal"
    assert styles["Given: F = 150 lb"] == "Callout"
    assert check(output).returncode == 0


def test_cover_and_info_bar_are_not_counted(tmp_path):
    md = "---\ntitle: Brakes\ncover: false\nName: [Name]\nClass: Auto 1\n---\nOne two three.\n"
    output, _ = build(tmp_path, md)
    assert "references): 3" in check(output).stdout


def test_header_and_footer_after_cover(tmp_path):
    output, _ = build(tmp_path, "---\ntitle: Brakes\nName: Jake Miller\n---\n# Intro\nText.\n")
    section = Document(str(output)).sections[0]
    assert section.different_first_page_header_footer
    assert section.header.paragraphs[0].text == "Brakes\tJake Miller"
    assert not section.first_page_header.paragraphs[0].text and not section.first_page_footer.paragraphs[0].text
    assert "NUMPAGES" in section.footer._element.xml
    assert not section.header.paragraphs[0].style.paragraph_format.tab_stops  # no stray center tab


def test_placeholder_name_is_not_put_in_the_header(tmp_path):
    output, _ = build(tmp_path, "---\ntitle: Brakes\nName: [Name]\nCourse: AUTO 110\n---\nText.\n")
    assert Document(str(output)).sections[0].header.paragraphs[0].text == "Brakes\tAUTO 110"


@pytest.mark.parametrize("theme", ["navy", "graphite", "plain"])
def test_themes_and_fonts(tmp_path, theme):
    md = f"---\ntitle: T\ntheme: {theme}\nfont: Arial\n---\n# Intro\n| A | B |\n|---|---|\n| x | 1 mm |\n"
    output, _ = build(tmp_path, md)
    doc = Document(str(output))
    fonts = doc.styles["Heading 1"].element.rPr.rFonts
    assert fonts.get(qn("w:ascii")) == "Arial" and fonts.get(qn("w:asciiTheme")) is None
    cell = doc.tables[-1].cell(1, 1).paragraphs[0]
    assert cell.alignment == 1  # numeric column centered


def test_unknown_theme_falls_back_to_navy(tmp_path):
    output, _ = build(tmp_path, "---\ntitle: T\ntheme: purple\n---\n# Intro\n")
    assert str(Document(str(output)).styles["Heading 1"].font.color.rgb) == "1F3864"


# ------------------------------------------------------------- picture placeholders

def test_picture_placeholder_box_with_hint_and_caption(tmp_path):
    md = ("---\ncover: false\n---\n# Job\nI swapped the pads.\n"
          '![Figure 1: Old pads](placeholder "Photo of the old pads next to the new ones"){width=10 height=6}\n'
          "![Figure 2: Torque wrench](photo)\nDone.\n")
    output, stdout = build(tmp_path, md)
    assert "2 picture placeholder(s)" in stdout and "Figure 1: Old pads" in stdout
    doc = Document(str(output))
    boxes = [t for t in doc.tables if t.cell(0, 0).paragraphs[0].style.name == "Picture Placeholder"]
    assert len(boxes) == 2 and len(doc.inline_shapes) == 0
    first = boxes[0].cell(0, 0)
    assert "INSERT PICTURE HERE" in first.text and "Photo of the old pads" in first.text
    assert 'w:val="dashed"' in first._tc.xml
    assert round(boxes[0].columns[0].width.cm) == 10 and round(boxes[0].rows[0].height.cm) == 6
    assert any(p.style.name == "Caption" and p.text == "Figure 1: Old pads" for p in doc.paragraphs)
    md_count = check(tmp_path / "answers.md").stdout.split("): ")[1].split()[0]
    stdout = check(output).stdout
    assert stdout.split("): ")[1].split()[0] == md_count  # box text is not counted
    assert "picture placeholders for the student to fill in: 2" in stdout


def test_missing_image_becomes_a_placeholder(tmp_path):
    output, stdout = build(tmp_path, "---\ncover: false\n---\n![Figure 3: Gauge reading](gauge.jpg)\n")
    assert "not found: gauge.jpg" in stdout and "placeholder box was put in its place" in stdout
    doc = Document(str(output))
    assert doc.tables and "INSERT PICTURE HERE" in doc.tables[-1].cell(0, 0).text


# ------------------------------------------------------------- review regressions

def test_dashes_in_detail_keys_and_font_never_reach_the_file(tmp_path):
    md = f"---\ntitle: T\nfont: Segoe UI {EN} Light\nCourse {EM} Section: AUTO 110\n---\nText.\n"
    output, stdout = build(tmp_path, md)
    assert not any(char in all_xml(output) for char in DASH_CHARS)
    assert check(output).returncode == 0


@pytest.mark.parametrize("name", ["Screenshot (1).png", "image (2).jpg"])
def test_image_path_with_brackets(tmp_path, name):
    from PIL import Image
    Image.new("RGB", (40, 20), "white").save(tmp_path / name)
    output, _ = build(tmp_path, f"---\ncover: false\n---\n![Figure 1: Pads]({name}){{width=8}}\n")
    doc = Document(str(output))
    assert len(doc.inline_shapes) == 1 and not any("![" in p.text for p in doc.paragraphs)


def test_tall_image_is_scaled_to_fit_the_page(tmp_path):
    from PIL import Image
    Image.new("RGB", (1170, 2532), "white").save(tmp_path / "phone.png")
    output, _ = build(tmp_path, "---\ncover: false\n---\n![Figure 1: Screen](phone.png)\n")
    shape = Document(str(output)).inline_shapes[0]
    assert shape.height.cm <= 19 and abs(shape.width / shape.height - 1170 / 2532) < 0.01


def test_consecutive_questions_stay_separate_and_blank_lines_are_kept(tmp_path):
    md = "---\ncover: false\n---\n**1. First question (1 pt)**\n**2. Second question (1 pt)**\n**3. Q three**\n___\n"
    output, _ = build(tmp_path, md)
    paragraphs = Document(str(output)).paragraphs
    second = next(p for p in paragraphs if p.text.startswith("2. Second"))
    assert second.paragraph_format.space_before is None  # keeps the style's gap, not merged into question 1
    assert any(p.text == "___" for p in paragraphs)


def test_toc_text_matches_heading_text(tmp_path):
    output, _ = build(tmp_path, "---\ntitle: T\ntoc: true\n---\n## 2 * 3 = 6\n# Fill ___ in\n", "--no-render")
    toc = [p.text for p in Document(str(output)).paragraphs if p.style.name.startswith("toc")]
    assert toc == ["2 * 3 = 6", "Fill ___ in"]


@needs_libreoffice
def test_toc_ignores_level_three_heading_with_same_text(tmp_path):
    md = "---\ntitle: T\ntoc: true\n---\n# Intro\n### Results\nText.\n\\pagebreak\n## Results\nMore.\n"
    output, _ = build(tmp_path, md)
    toc = [p.text for p in Document(str(output)).paragraphs if p.style.name.startswith("toc")]
    assert toc == ["Intro\t3", "Results\t4"]


def test_phone_photo_exif_rotation_is_applied(tmp_path):
    from PIL import Image
    image = Image.new("RGB", (400, 200), "red")
    exif = image.getexif()
    exif[0x0112] = 6  # rotate 90 degrees when shown
    image.save(tmp_path / "photo.jpg", exif=exif.tobytes())
    output, _ = build(tmp_path, "---\ncover: false\n---\n![Figure 1: Pads](photo.jpg){width=5}\n")
    shape = Document(str(output)).inline_shapes[0]
    assert shape.height > shape.width  # shown upright (portrait), like on the phone


def test_long_title_and_control_characters_do_not_crash(tmp_path):
    md = "---\ntitle: " + "Brakes " * 50 + "\nsubtitle: " + "x" * 400 + "\n---\nText\x00 with\x01 control\x08 chars.\n"
    output, _ = build(tmp_path, md)
    assert any("Text with control chars." == p.text for p in Document(str(output)).paragraphs)


def test_bold_with_italic_inside(tmp_path):
    output, _ = build(tmp_path, "---\ncover: false\n---\n**Note: this is *really* important**\n")
    paragraph = next(p for p in Document(str(output)).paragraphs if "Note" in p.text)
    assert paragraph.text == "Note: this is really important"
    assert all(run.bold for run in paragraph.runs if run.text.strip())
    assert any(run.italic for run in paragraph.runs if run.text == "really")


def test_question_detection_false_positives(tmp_path):
    md = ("---\ncover: false\n---\n**5. Calculation (4 pts)**\n**a) Find the speed at point B. Use v = d/t**\nAnswer.\n"
          "**3.5 L engine specs**\n**12.5 mm minimum**\n**2.1 Explain what a thermostat does.**\n")
    output, _ = build(tmp_path, md)
    styles = {p.text: p.style.name for p in Document(str(output)).paragraphs if p.text}
    assert styles["a) Find the speed at point B. Use v = d/t"] == "Sub Question"
    assert styles["3.5 L engine specs"] == "Normal" and styles["12.5 mm minimum"] == "Normal"
    assert styles["2.1 Explain what a thermostat does."] == "Question"


def test_numbered_item_keeps_its_tab_and_url_keeps_double_hyphen(tmp_path):
    output, _ = build(tmp_path, f"---\ncover: false\n---\n1. {EM} see above\nLink https://x.com/a--b here\n")
    texts = [p.text for p in Document(str(output)).paragraphs]
    assert "1.\tsee above" in texts and "Link https://x.com/a--b here" in texts


def test_small_fixes_size_lang_tabs_pagebreak(tmp_path):
    md = ("---\ntitle: T\ncover: false\nlang: en_GB\nName: **Jake**\n---\n- top\n\t- nested\n"
          "![Fig](missing.png){width=1.2.3}\n[pagebreak]\nOne two.\n")
    output, _ = build(tmp_path, md)
    doc = Document(str(output))
    assert round(doc.sections[0].page_height.cm, 1) == 29.7
    assert any(p.style.name == "List Bullet 2" for p in doc.paragraphs)
    assert doc.sections[0].header.paragraphs[0].text == "T\tJake"
    md_count = check(tmp_path / "answers.md").stdout.split("): ")[1].split()[0]
    assert check(output).stdout.split("): ")[1].split()[0] == md_count


def test_metadata_time_is_utc_and_dash_count_is_not_inflated(tmp_path):
    import datetime as dt
    output, stdout = build(tmp_path, f"---\ntitle: Brakes {EM} test\nName: Jake\n---\nText.\n")
    created = Document(str(output)).core_properties.created.replace(tzinfo=dt.timezone.utc)
    assert abs((dt.datetime.now(dt.timezone.utc) - created).total_seconds()) < 120
    assert "removed 1 em/en dash" in stdout


def test_failed_image_leaves_no_empty_paragraph(tmp_path):
    (tmp_path / "broken.png").write_bytes(b"not an image")
    output, stdout = build(tmp_path, "---\ncover: false\n---\nBefore.\n![Fig](broken.png)\n")
    body = Document(str(output)).element.body
    empty = [p for p in body.iterchildren() if p.tag.endswith("}p") and not "".join(p.itertext()).strip()
             and "lineRule" not in p.xml and "sectPr" not in p.xml]
    assert not empty and "unreadable: broken.png ()" not in stdout


def test_failed_conversion_never_returns_a_stale_pdf(tmp_path, monkeypatch):
    monkeypatch.setattr(office_render, "sandbox_runner", lambda: None)
    monkeypatch.setattr(office_render, "find_soffice", lambda: "/usr/bin/true")  # "converts" nothing
    (tmp_path / "doc.docx").write_bytes(b"x")
    stale = tmp_path / "out" / "doc.pdf"
    stale.parent.mkdir()
    stale.write_bytes(b"%PDF-1.4 old render")
    assert office_render.to_pdf(tmp_path / "doc.docx", tmp_path / "out") is None and not stale.exists()
